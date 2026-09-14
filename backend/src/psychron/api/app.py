"""The read API.

Every route is authenticated except the health probe, and every range query is
bounded before it reaches the database — see adapters/queries.py for why the
bucket is chosen from the span rather than asked for by the caller.
"""

from __future__ import annotations

import asyncio
import csv
import io
import json
import logging
from datetime import datetime, timedelta, timezone
from typing import Iterator

from fastapi import Depends, FastAPI, HTTPException, Query, WebSocket, WebSocketDisconnect
from fastapi.responses import StreamingResponse

from ..adapters.identity_pg import PostgresIdentity
from ..adapters.queries import MAX_POINTS, ReadQueries
from ..domain import psychrometrics
from .auth import (Principal, TokenAuthenticator, ensure_token, require_principal,
                   set_authenticator, set_identity)
from .routes_auth import SESSION_COOKIE, build_router

log = logging.getLogger(__name__)

# Quality bits, so an export explains itself instead of shipping a bare integer.
QUALITY_FLAGS = {
    0x01: "clock_unsynced", 0x02: "replayed", 0x04: "prior_read_failed",
    0x08: "out_of_range", 0x10: "interval_drift",
    0x0100: "time_from_boot_anchor", 0x0200: "time_from_arrival",
}


def describe_quality(value: int | None) -> list[str]:
    if not value:
        return []
    return [name for bit, name in QUALITY_FLAGS.items() if value & bit]


def parse_range(frm: str | None, to: str | None) -> tuple[datetime, datetime]:
    now = datetime.now(timezone.utc)
    try:
        end = datetime.fromisoformat(to) if to else now
        start = datetime.fromisoformat(frm) if frm else end - timedelta(hours=24)
    except ValueError as exc:
        raise HTTPException(400, f"bad timestamp: {exc}") from exc

    if start.tzinfo is None:
        start = start.replace(tzinfo=timezone.utc)
    if end.tzinfo is None:
        end = end.replace(tzinfo=timezone.utc)
    if start >= end:
        raise HTTPException(400, "'from' must be before 'to'")
    if end - start > timedelta(days=3650):
        raise HTTPException(400, "range longer than ten years")
    return start, end


def create_app(dsn: str, env_path, device_id: str = "esp32-01",
               phone_device_id: str = "phone-01") -> FastAPI:
    set_authenticator(TokenAuthenticator(ensure_token(env_path)))
    identity = PostgresIdentity(dsn)
    set_identity(identity)
    queries = ReadQueries(dsn)

    app = FastAPI(title="psychron", version="1", docs_url="/api/docs")
    app.include_router(build_router(identity))

    # No CORS middleware. The dashboard is served from the same origin as the
    # API — by a Vite proxy in development and by Caddy in production — so there
    # is no cross-origin request to permit. An allowlist that is never exercised
    # is an allowlist that drifts out of date and is eventually widened by
    # someone debugging a problem it was not causing.

    @app.get("/api/health")
    def health() -> dict:
        """Unauthenticated on purpose: a liveness probe that needs a secret is
        useless to the thing most likely to be checking it."""
        latest = queries.latest(device_id)
        age = None
        if latest:
            age = (datetime.now(timezone.utc) - latest["time"]).total_seconds()
        phone = queries.sample_latest(phone_device_id)
        phone_age = None
        if phone:
            phone_age = (datetime.now(timezone.utc) - phone["time"]).total_seconds()
        return {
            "status": "ok",
            "device": device_id,
            "last_reading_age_s": age,
            # A node is not "up" because the service is: say so separately.
            "device_reporting": age is not None and age < 30,
            "phone": phone_device_id,
            "last_sample_age_s": phone_age,
            "phone_reporting": phone_age is not None and phone_age < 30,
        }

    @app.get("/api/current")
    def current(_: Principal = Depends(require_principal)) -> dict:
        row = queries.latest(device_id)
        if row is None:
            raise HTTPException(404, "no readings stored yet")

        derived = psychrometrics.derive(row["temperature_c"], row["humidity_pct"])

        # The same room a day earlier, so "23.4 °C" can be read as a movement
        # rather than a bare number. None when nothing was recorded near that
        # instant — an outage must not be papered over with a value from
        # whenever the device happened to be alive.
        day_ago = queries.near(device_id, datetime.now(timezone.utc) - timedelta(days=1))

        return {
            # Measured and derived are kept apart deliberately: one carries
            # sensor error, the other sensor error plus a fitted equation's.
            "measured": {
                "time": row["time"],
                "temperature_c": row["temperature_c"],
                "humidity_pct": row["humidity_pct"],
            },
            "derived": derived,
            "provenance": {
                "boot_id": row["boot_id"], "seq": row["seq"],
                "firmware": row["firmware"],
                "device_time": row["device_time"],
                "received_at": row["received_at"],
                "quality": row["quality"],
                "quality_flags": describe_quality(row["quality"]),
            },
            "day_ago": day_ago,
        }

    @app.get("/api/readings")
    def readings(frm: str | None = Query(None, alias="from"),
                 to: str | None = None,
                 _: Principal = Depends(require_principal)) -> dict:
        start, end = parse_range(frm, to)
        series = queries.series(device_id, start, end)
        return {
            "from": start, "to": end,
            "bucket": series.bucket,
            "count": len(series.points),
            # Truncation is reported, never silent: a chart missing its tail
            # looks exactly like a sensor that stopped.
            "truncated": series.truncated,
            "max_points": MAX_POINTS,
            "points": series.points,
        }

    @app.get("/api/stats")
    def stats(frm: str | None = Query(None, alias="from"),
              to: str | None = None,
              _: Principal = Depends(require_principal)) -> dict:
        start, end = parse_range(frm, to)
        return {"from": start, "to": end, **queries.stats(device_id, start, end)}

    @app.get("/api/device")
    def device(_: Principal = Depends(require_principal)) -> dict:
        return {"device": device_id,
                "extent": queries.extent(device_id),
                **queries.device_health(device_id)}

    # ── export ──────────────────────────────────────────────────────────────

    def export_rows(start: datetime, end: datetime) -> Iterator[dict]:
        for row in queries.stream_readings(device_id, start, end):
            derived = psychrometrics.derive(row["temperature_c"], row["humidity_pct"])
            yield {**row, **derived,
                   "quality_flags": "|".join(describe_quality(row["quality"]))}

    COLUMNS = ["time", "device_id", "boot_id", "seq", "device_time", "received_at",
               "uptime_ms", "temperature_c", "humidity_pct", "dew_point_c",
               "absolute_humidity_g_m3", "vapour_pressure_deficit_kpa",
               "heat_index_c", "condensation_margin_c", "quality", "quality_flags",
               "firmware"]

    UNITS = {"temperature_c": "degC", "humidity_pct": "percent", "dew_point_c": "degC",
             "absolute_humidity_g_m3": "g/m3", "vapour_pressure_deficit_kpa": "kPa",
             "heat_index_c": "degC", "condensation_margin_c": "degC",
             "uptime_ms": "ms"}

    def csv_stream(start, end) -> Iterator[str]:
        buf = io.StringIO()
        writer = csv.writer(buf, lineterminator="\n")

        # Units and provenance travel with the data. A column of bare numbers
        # is a column somebody will later assume is Fahrenheit.
        writer.writerow([f"# psychron export  device={device_id}  "
                         f"from={start.isoformat()}  to={end.isoformat()}"])
        writer.writerow(["# units: " + ", ".join(f"{k}={v}" for k, v in UNITS.items())])
        writer.writerow(COLUMNS)
        yield buf.getvalue(); buf.seek(0); buf.truncate()

        for row in export_rows(start, end):
            writer.writerow([row.get(c) for c in COLUMNS])
            yield buf.getvalue(); buf.seek(0); buf.truncate()

    def json_stream(start, end) -> Iterator[str]:
        head = {"device": device_id, "from": start.isoformat(), "to": end.isoformat(),
                "units": UNITS, "quality_flags": QUALITY_FLAGS}
        yield '{"metadata": ' + json.dumps(head) + ', "readings": ['
        first = True
        for row in export_rows(start, end):
            yield ("" if first else ",") + json.dumps(row, default=str)
            first = False
        yield "]}"

    def xlsx_bytes(start, end) -> bytes:
        from openpyxl import Workbook
        wb = Workbook(write_only=True)
        ws = wb.create_sheet("readings")
        ws.append(COLUMNS)
        for row in export_rows(start, end):
            ws.append([v.isoformat() if isinstance(v, datetime) else v
                       for v in (row.get(c) for c in COLUMNS)])
        meta = wb.create_sheet("metadata")
        meta.append(["device", device_id])
        meta.append(["from", start.isoformat()])
        meta.append(["to", end.isoformat()])
        for k, v in UNITS.items():
            meta.append([k, v])
        out = io.BytesIO()
        wb.save(out)
        return out.getvalue()

    @app.get("/api/export")
    def export(fmt: str = Query("csv", alias="format"),
               frm: str | None = Query(None, alias="from"),
               to: str | None = None,
               _: Principal = Depends(require_principal)):
        start, end = parse_range(frm, to)
        stamp = f"{start:%Y%m%d}-{end:%Y%m%d}"

        if fmt == "csv":
            return StreamingResponse(
                csv_stream(start, end), media_type="text/csv",
                headers={"Content-Disposition":
                         f'attachment; filename="psychron-{stamp}.csv"'})
        if fmt == "json":
            return StreamingResponse(
                json_stream(start, end), media_type="application/json",
                headers={"Content-Disposition":
                         f'attachment; filename="psychron-{stamp}.json"'})
        if fmt in ("xlsx", "excel"):
            # Not streamable: the format is a zip archive with a central
            # directory, so it exists only once complete. Bounded instead.
            n = queries.count_readings(device_id, start, end)
            if n > 200_000:
                raise HTTPException(
                    413, f"{n:,} rows is too many for a spreadsheet; "
                         "narrow the range or export CSV, which streams")
            return StreamingResponse(
                iter([xlsx_bytes(start, end)]),
                media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                headers={"Content-Disposition":
                         f'attachment; filename="psychron-{stamp}.xlsx"'})
        raise HTTPException(400, "format must be csv, json or xlsx")

    @app.get("/api/export/preview")
    def export_preview(frm: str | None = Query(None, alias="from"),
                       to: str | None = None,
                       _: Principal = Depends(require_principal)) -> dict:
        """What a download would contain, before committing to it."""
        start, end = parse_range(frm, to)
        rows = queries.count_readings(device_id, start, end)
        return {"from": start, "to": end, "rows": rows,
                "columns": COLUMNS,
                "estimated_csv_bytes": rows * 180,
                "xlsx_available": rows <= 200_000}

    # ── live ────────────────────────────────────────────────────────────────

    def _socket_allowed(ws: WebSocket) -> bool:
        # Cookies travel with the WebSocket handshake, so a browser needs no
        # token in the URL — which is where the old scheme leaked it into logs
        # and history. The query parameter remains only for scripted clients.
        # One function for every socket: two copies of an admission check are
        # two chances for one of them to be wrong.
        cookie = ws.cookies.get(SESSION_COOKIE)
        if cookie and identity.session_principal(cookie) is not None:
            return True
        from .auth import _authenticator
        token = ws.query_params.get("token", "")
        return (_authenticator is not None
                and _authenticator.authenticate(f"Bearer {token}") is not None)

    def _utc(instant: datetime) -> str:
        # The same spelling FastAPI uses for the ESP32's readings. "+00:00" and "Z"
        # are the same instant, but two payloads from one system should not look
        # as if they came from two clocks when someone lays them side by side.
        return instant.astimezone(timezone.utc).isoformat().replace("+00:00", "Z")

    def _sample_payload(row: dict) -> dict:
        out = {"time": _utc(row["time"]),
               "window_ms": row["window_ms"],
               "firmware": row["firmware"],
               "quality_flags": describe_quality(row["quality"])}
        for column in ReadQueries.SAMPLE_COLUMNS:
            out[column] = row[column]
        return out

    @app.get("/api/phone/current")
    def phone_current(_: Principal = Depends(require_principal)) -> dict:
        row = queries.sample_latest(phone_device_id)
        if row is None:
            raise HTTPException(404, "no samples stored yet")
        tendency = None
        if row["pressure_hpa"] is not None:
            tendency = queries.pressure_tendency(phone_device_id, row["pressure_hpa"], row["time"])
        return {"device": phone_device_id,
                **_sample_payload(row),
                "pressure_tendency_3h_hpa": tendency}

    @app.get("/api/phone/series")
    def phone_series(frm: str | None = Query(None, alias="from"),
                     to: str | None = None,
                     _: Principal = Depends(require_principal)) -> dict:
        start, end = parse_range(frm, to)
        series = queries.sample_series(phone_device_id, start, end)
        return {"from": start, "to": end, "bucket": series.bucket,
                "count": len(series.points), "truncated": series.truncated,
                "max_points": MAX_POINTS, "points": series.points}

    @app.websocket("/api/live/phone")
    async def live_phone(ws: WebSocket) -> None:
        if not _socket_allowed(ws):
            await ws.close(code=4401)
            return
        await ws.accept()
        last_seen: datetime | None = None
        try:
            while True:
                row = await asyncio.to_thread(queries.sample_latest, phone_device_id)
                if row and row["time"] != last_seen:
                    last_seen = row["time"]
                    await ws.send_json(_sample_payload(row))
                # Half the node's window: often enough that a gesture in front of
                # the phone shows up without a visible lag, rarely enough that an
                # idle panel is not a query storm.
                await asyncio.sleep(1.0)
        except WebSocketDisconnect:
            pass

    @app.websocket("/api/live")
    async def live(ws: WebSocket) -> None:
        if not _socket_allowed(ws):
            await ws.close(code=4401)
            return

        await ws.accept()
        last_seen: datetime | None = None
        try:
            while True:
                row = await asyncio.to_thread(queries.latest, device_id)
                if row and row["time"] != last_seen:
                    last_seen = row["time"]
                    await ws.send_json({
                        "time": row["time"].isoformat(),
                        "temperature_c": row["temperature_c"],
                        "humidity_pct": row["humidity_pct"],
                        "quality_flags": describe_quality(row["quality"]),
                        **{k: v for k, v in psychrometrics.derive(
                            row["temperature_c"], row["humidity_pct"]).items()},
                    })
                # Polling, not LISTEN/NOTIFY. One device at three seconds does
                # not justify a trigger on the hot write path; this changes if a
                # fleet ever arrives.
                await asyncio.sleep(1.0)
        except WebSocketDisconnect:
            pass

    return app
